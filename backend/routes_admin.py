import io
import csv
import os
from datetime import timedelta

from fastapi import APIRouter, Depends, HTTPException, Request, Response, UploadFile, File, Form
from pydantic import BaseModel, Field, EmailStr
from typing import Optional, List, Literal

from core import (db, now, now_iso, new_id, slugify, hash_password, verify_password, create_token, current_user,
                  require, has_perm, actor_of, audit, put_object, get_object, decrypt, APP_NAME, IMAGE_TYPES,
                  client_ip, PERMS)
from catalog import recompute, lookup_maps
import orders as O
import tg_api
import telegram_bot as TB

router = APIRouter(prefix="/api")


# ---------------- auth ----------------
class LoginIn(BaseModel):
    email: EmailStr
    password: str = Field(min_length=1, max_length=200)


def set_cookies(resp, uid):
    resp.set_cookie("access_token", create_token(uid, "access", 60), httponly=True, secure=True, samesite="none", max_age=3600, path="/")
    resp.set_cookie("refresh_token", create_token(uid, "refresh", 60 * 24 * 7), httponly=True, secure=True, samesite="none", max_age=604800, path="/")


def user_out(u):
    u = {k: v for k, v in u.items() if k not in ("password_hash", "_id")}
    u["permissions"] = sorted(PERMS.get(u["role"], set()) | ({"orders.approve"} if u.get("can_approve_payments") else set()))
    return u


@router.post("/auth/login")
async def login(body: LoginIn, request: Request, response: Response):
    email = body.email.lower()
    ident = f"{client_ip(request)}:{email}"
    att = await db.login_attempts.find_one({"identifier": ident}, {"_id": 0})
    if att and att.get("count", 0) >= 5:
        if att.get("locked_until", "") > now_iso():
            raise HTTPException(429, "Too many failed attempts. Try again in 15 minutes.")
        att = None
        await db.login_attempts.delete_one({"identifier": ident})
    u = await db.users.find_one({"email": email, "deleted_at": None}, {"_id": 0})
    if not u or not u.get("active") or not verify_password(body.password, u["password_hash"]):
        cnt = (att or {}).get("count", 0) + 1
        await db.login_attempts.update_one({"identifier": ident}, {"$set": {"count": cnt, "locked_until": (now() + timedelta(minutes=15)).isoformat()}}, upsert=True)
        raise HTTPException(401, "Invalid email or password")
    await db.login_attempts.delete_one({"identifier": ident})
    set_cookies(response, u["id"])
    await audit({"type": "admin_user", "id": u["id"], "name": u["email"]}, "auth.login", "user", u["id"])
    return {**user_out(u), "access_token": create_token(u["id"], "access", 60)}


@router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/", secure=True, samesite="none")
    response.delete_cookie("refresh_token", path="/", secure=True, samesite="none")
    return {"ok": True}


@router.get("/auth/me")
async def me(user=Depends(current_user)):
    return user_out(user)


@router.post("/auth/refresh")
async def refresh(request: Request, response: Response):
    import jwt as _jwt
    tok = request.cookies.get("refresh_token")
    try:
        payload = _jwt.decode(tok or "", os.environ["JWT_SECRET"], algorithms=["HS256"])
        assert payload.get("type") == "refresh"
    except Exception:
        raise HTTPException(401, "Invalid refresh token")
    if not await db.users.find_one({"id": payload["sub"], "active": True, "deleted_at": None}):
        raise HTTPException(401, "User inactive")
    set_cookies(response, payload["sub"])
    return {"ok": True}


class PasswordIn(BaseModel):
    current_password: str
    new_password: str = Field(min_length=8, max_length=200)


@router.post("/auth/change-password")
async def change_password(body: PasswordIn, user=Depends(current_user)):
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0})
    if not verify_password(body.current_password, u["password_hash"]):
        raise HTTPException(400, "Current password is incorrect")
    await db.users.update_one({"id": user["id"]}, {"$set": {"password_hash": hash_password(body.new_password)}})
    await audit(actor_of(user), "auth.password_changed", "user", user["id"])
    return {"ok": True}


# ---------------- users ----------------
class UserIn(BaseModel):
    email: EmailStr
    name: str = Field(min_length=1, max_length=80)
    role: Literal["owner", "admin", "product_manager", "customer_service"]
    password: Optional[str] = Field(default=None, min_length=8, max_length=200)
    can_approve_payments: bool = False
    active: bool = True


@router.get("/admin/users")
async def users(user=Depends(require("users.view"))):
    return [user_out(u) for u in await db.users.find({"deleted_at": None}, {"_id": 0}).to_list(200)]


@router.post("/admin/users")
async def create_user(body: UserIn, user=Depends(require("users"))):
    if not body.password:
        raise HTTPException(400, "Password required")
    if await db.users.find_one({"email": body.email.lower()}):
        raise HTTPException(400, "Email already exists")
    u = {"id": new_id(), "email": body.email.lower(), "name": body.name, "role": body.role, "active": body.active,
         "can_approve_payments": body.can_approve_payments, "password_hash": hash_password(body.password),
         "created_at": now_iso(), "deleted_at": None}
    await db.users.insert_one(dict(u))
    await audit(actor_of(user), "user.created", "user", u["id"], {"email": u["email"], "role": u["role"]})
    return user_out(u)


@router.put("/admin/users/{uid}")
async def update_user(uid: str, body: UserIn, user=Depends(require("users"))):
    target = await db.users.find_one({"id": uid, "deleted_at": None}, {"_id": 0})
    if not target:
        raise HTTPException(404, "User not found")
    if target["role"] == "owner" and (body.role != "owner" or not body.active):
        if await db.users.count_documents({"role": "owner", "active": True, "deleted_at": None}) <= 1:
            raise HTTPException(400, "Cannot demote or deactivate the last owner")
    upd = {"name": body.name, "role": body.role, "active": body.active, "can_approve_payments": body.can_approve_payments}
    if body.password:
        upd["password_hash"] = hash_password(body.password)
    await db.users.update_one({"id": uid}, {"$set": upd})
    await audit(actor_of(user), "user.updated", "user", uid, {"role": body.role, "active": body.active, "password_reset": bool(body.password)})
    return user_out(await db.users.find_one({"id": uid}, {"_id": 0}))


# ---------------- brands / categories ----------------
class BrandIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)
    logo_url: Optional[str] = ""
    color: Optional[str] = "#2563EB"
    sort: int = 100


class CategoryIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)
    name_my: Optional[str] = ""
    icon: Optional[str] = "sparkles"
    sort: int = 100


def make_crud(coll, model, label):
    @router.get(f"/admin/{coll}")
    async def _list(user=Depends(require("catalog"))):
        return await db[coll].find({"deleted_at": None}, {"_id": 0}).sort("sort", 1).to_list(500)

    @router.post(f"/admin/{coll}")
    async def _create(body: model, user=Depends(require("catalog"))):
        d = {"id": new_id(), **body.model_dump(), "slug": slugify(body.name).rsplit("-", 1)[0], "created_at": now_iso(), "deleted_at": None}
        if await db[coll].find_one({"slug": d["slug"], "deleted_at": None}):
            raise HTTPException(400, f"{label} already exists")
        await db[coll].insert_one(dict(d))
        await audit(actor_of(user), f"{label}.created", label, d["id"], {"name": body.name})
        return d

    @router.put(f"/admin/{coll}/{{iid}}")
    async def _update(iid: str, body: model, user=Depends(require("catalog"))):
        r = await db[coll].update_one({"id": iid, "deleted_at": None}, {"$set": body.model_dump()})
        if not r.matched_count:
            raise HTTPException(404, "Not found")
        await audit(actor_of(user), f"{label}.updated", label, iid, body.model_dump())
        return await db[coll].find_one({"id": iid}, {"_id": 0})

    @router.delete(f"/admin/{coll}/{{iid}}")
    async def _delete(iid: str, user=Depends(require("catalog"))):
        await db[coll].update_one({"id": iid}, {"$set": {"deleted_at": now_iso()}})
        await audit(actor_of(user), f"{label}.deleted", label, iid)
        return {"ok": True}


make_crud("brands", BrandIn, "brand")
make_crud("categories", CategoryIn, "category")


# ---------------- products ----------------
class VariantIn(BaseModel):
    id: Optional[str] = None
    name: str = Field(min_length=1, max_length=80)
    duration: Optional[str] = Field(default="", max_length=40)
    price_mmk: int = Field(ge=0, le=100_000_000)
    promo_price_mmk: Optional[int] = Field(default=None, ge=0)
    promo_start: Optional[str] = None
    promo_end: Optional[str] = None
    stock_status: Literal["in_stock", "out_of_stock", "limited"] = "in_stock"
    stock_qty: Optional[int] = Field(default=None, ge=0)
    stock_override: Optional[Literal["in_stock", "out_of_stock"]] = None
    reseller_price_mmk: Optional[int] = Field(default=None, ge=0, le=100_000_000)
    supplier_id: Optional[str] = None
    supplier_product_id: Optional[str] = None
    price_locked: bool = True
    active: bool = True


class ProductIn(BaseModel):
    name: str = Field(min_length=2, max_length=160)
    name_my: Optional[str] = ""
    brand_id: Optional[str] = None
    category_id: Optional[str] = None
    short_description: Optional[str] = Field(default="", max_length=300)
    short_description_my: Optional[str] = Field(default="", max_length=600)
    description: Optional[str] = Field(default="", max_length=8000)
    description_my: Optional[str] = Field(default="", max_length=12000)
    product_type: Optional[str] = "subscription"
    activation_method: Optional[str] = ""
    eligibility: Optional[str] = ""
    warranty: Optional[str] = ""
    delivery_time: Optional[str] = ""
    images: List[str] = []
    featured: bool = False
    hot: bool = False
    status: Literal["draft", "published", "unpublished", "archived"] = "draft"
    variants: List[VariantIn] = Field(min_length=1)


SUPPLIER_OWNED = ("supplier_cost", "supplier_stock", "last_synced_at", "supplier_variant_key", "supplier_currency")


async def build_product(body: ProductIn, existing=None, actor=None):
    for v in body.variants:
        if v.promo_price_mmk is not None and v.promo_price_mmk >= v.price_mmk:
            raise HTTPException(400, f"Promo price for '{v.name}' must be lower than the regular price")
    old = {v["id"]: v for v in (existing or {}).get("variants", [])}
    variants = []
    for v in body.variants:
        d = v.model_dump()
        d["id"] = d["id"] if d["id"] in old else new_id()
        prev = old.get(d["id"], {})
        for k in SUPPLIER_OWNED:
            d[k] = prev.get(k)
        if prev and prev.get("price_mmk") != d["price_mmk"]:
            await db.price_history.insert_one({"id": new_id(), "kind": "selling_price", "product_id": existing["id"], "variant_id": d["id"],
                                               "old": prev.get("price_mmk"), "new": d["price_mmk"], "source": "admin", "actor": actor, "created_at": now_iso()})
        if prev and prev.get("stock_override") != d["stock_override"]:
            await db.stock_overrides.insert_one({"id": new_id(), "product_id": existing["id"], "variant_id": d["id"],
                                                 "old": prev.get("stock_override"), "new": d["stock_override"], "actor": actor, "created_at": now_iso()})
        variants.append(d)
    data = body.model_dump()
    data["variants"] = variants
    return recompute(data)


@router.get("/admin/products")
async def admin_products(q: str = "", status: str = "", source: str = "", user=Depends(require("products"))):
    import re
    query = {"deleted_at": None}
    if q:
        query["name"] = {"$regex": re.escape(q), "$options": "i"}
    if status:
        query["status"] = status
    if source:
        query["source_type"] = source
    items = await db.products.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    bm, cm = await lookup_maps()
    for p in items:
        p["brand_name"] = (bm.get(p.get("brand_id")) or {}).get("name")
        p["category_name"] = (cm.get(p.get("category_id")) or {}).get("name")
    return items


@router.get("/admin/products/{pid}")
async def admin_product(pid: str, user=Depends(require("products"))):
    p = await db.products.find_one({"id": pid, "deleted_at": None}, {"_id": 0})
    if not p:
        raise HTTPException(404, "Product not found")
    p["price_history"] = await db.price_history.find({"product_id": pid}, {"_id": 0}).sort("created_at", -1).limit(20).to_list(20)
    return p


@router.post("/admin/products")
async def create_product(body: ProductIn, user=Depends(require("products"))):
    d = await build_product(body, actor=actor_of(user))
    d.update(id=new_id(), slug=slugify(body.name), source_type="MANUAL", order_count=0, created_at=now_iso(),
             updated_at=now_iso(), deleted_at=None, is_sample=False)
    await db.products.insert_one(dict(d))
    await audit(actor_of(user), "product.created", "product", d["id"], {"name": d["name"], "status": d["status"]})
    return d


@router.put("/admin/products/{pid}")
async def update_product(pid: str, body: ProductIn, user=Depends(require("products"))):
    ex = await db.products.find_one({"id": pid, "deleted_at": None}, {"_id": 0})
    if not ex:
        raise HTTPException(404, "Product not found")
    d = await build_product(body, ex, actor_of(user))
    d.update(updated_at=now_iso(), is_sample=ex.get("is_sample", False) and d["name"] == ex["name"])
    await db.products.update_one({"id": pid}, {"$set": d})
    await audit(actor_of(user), "product.updated", "product", pid, {"name": d["name"], "status": d["status"]})
    return await db.products.find_one({"id": pid}, {"_id": 0})


@router.delete("/admin/products/{pid}")
async def delete_product(pid: str, user=Depends(require("products"))):
    await db.products.update_one({"id": pid}, {"$set": {"deleted_at": now_iso(), "status": "archived"}})
    await audit(actor_of(user), "product.deleted", "product", pid)
    return {"ok": True}


class BulkIn(BaseModel):
    ids: List[str] = Field(min_length=1, max_length=500)
    action: Literal["publish", "unpublish", "archive", "feature", "unfeature", "hot", "unhot", "set_category", "set_brand", "price_adjust_percent"]
    value: Optional[str] = None


@router.post("/admin/products/bulk")
async def bulk(body: BulkIn, user=Depends(require("products"))):
    simple = {"publish": {"status": "published"}, "unpublish": {"status": "unpublished"}, "archive": {"status": "archived"},
              "feature": {"featured": True}, "unfeature": {"featured": False}, "hot": {"hot": True}, "unhot": {"hot": False},
              "set_category": {"category_id": body.value}, "set_brand": {"brand_id": body.value}}
    q = {"id": {"$in": body.ids}, "deleted_at": None}
    if body.action in simple:
        r = await db.products.update_many(q, {"$set": {**simple[body.action], "updated_at": now_iso()}})
        n = r.modified_count
    else:
        try:
            pct = float(body.value)
            assert -90 <= pct <= 500
        except Exception:
            raise HTTPException(400, "Percent must be a number between -90 and 500")
        n = 0
        async for p in db.products.find(q, {"_id": 0}):
            for v in p["variants"]:
                newp = int(round(v["price_mmk"] * (1 + pct / 100) / 100) * 100)
                await db.price_history.insert_one({"id": new_id(), "kind": "selling_price", "product_id": p["id"], "variant_id": v["id"],
                                                   "old": v["price_mmk"], "new": newp, "source": "bulk", "created_at": now_iso()})
                v["price_mmk"] = newp
                if v.get("promo_price_mmk") and v["promo_price_mmk"] >= newp:
                    v["promo_price_mmk"] = None
            recompute(p)
            await db.products.update_one({"id": p["id"]}, {"$set": p})
            n += 1
    await audit(actor_of(user), "product.bulk_" + body.action, "product", None, {"ids": body.ids, "value": body.value})
    return {"updated": n}


def parse_sheet(filename, raw):
    try:
        return _parse_sheet(filename, raw)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, "Could not read file. Use a UTF-8 .csv or a valid .xlsx") from e


def _parse_sheet(filename, raw):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        it = load_workbook(io.BytesIO(raw), read_only=True).active.iter_rows(values_only=True)
        hdr = [str(h or "").strip() for h in next(it)]
        return [{hdr[i]: (c if c is not None else "") for i, c in enumerate(r) if i < len(hdr)} for r in it]
    raise HTTPException(400, "Upload a .csv or .xlsx file")


def group_rows(rows):
    groups, errors = {}, []
    for i, r in enumerate(rows[:2000], 2):
        r = {str(k).strip().lower(): str(v).strip() for k, v in r.items()}
        if not r.get("name"):
            errors.append({"row": i, "error": "name is required"})
            continue
        try:
            r["_price"] = int(float(r.get("price_mmk") or "x"))
        except ValueError:
            errors.append({"row": i, "error": "price_mmk must be a number"})
            continue
        groups.setdefault(r["name"].lower(), []).append(r)
    return groups, errors


async def find_or_create(coll, name, extra):
    import re
    if not name:
        return None
    doc = await db[coll].find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}, "deleted_at": None})
    if not doc:
        doc = {"id": new_id(), "name": name, "slug": slugify(name).rsplit("-", 1)[0], "sort": 100,
               "created_at": now_iso(), "deleted_at": None, **extra}
        await db[coll].insert_one(dict(doc))
    return doc["id"]


def row_variant(x):
    stock = x.get("stock_status") if x.get("stock_status") in ("in_stock", "out_of_stock", "limited") else "in_stock"
    return VariantIn(name=x.get("plan") or "Standard", duration=x.get("duration", ""), price_mmk=x["_price"], stock_status=stock)


@router.post("/admin/products/import")
async def import_products(file: UploadFile = File(...), user=Depends(require("products"))):
    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(400, "File too large (max 5MB)")
    groups, errors = group_rows(parse_sheet(file.filename, raw))
    created = 0
    for rs in groups.values():
        r = rs[0]
        try:
            bid = await find_or_create("brands", r.get("brand"), {"color": "#2563EB"})
            cid = await find_or_create("categories", r.get("category"), {"name_my": "", "icon": "sparkles"})
            body = ProductIn(name=r["name"], brand_id=bid, category_id=cid, description=r.get("description", ""),
                             short_description=r.get("short_description", "")[:300], status="draft",
                             variants=[row_variant(x) for x in rs])
            await create_product(body, user)
            created += 1
        except Exception as e:
            errors.append({"row": r["name"], "error": str(getattr(e, "detail", e))[:200]})
    await audit(actor_of(user), "product.import", "product", None, {"created": created, "errors": len(errors)})
    return {"created": created, "errors": errors[:50]}


@router.post("/admin/uploads/image")
async def upload_image(file: UploadFile = File(...), user=Depends(require("products"))):
    ct = file.content_type or ""
    if ct not in IMAGE_TYPES:
        raise HTTPException(400, "Only JPG, PNG, WEBP or GIF images are allowed")
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(400, "Image too large (max 5MB)")
    res = await put_object(f"{APP_NAME}/public/images/{new_id()}.{IMAGE_TYPES[ct]}", data, ct)
    await db.files.insert_one({"id": new_id(), "storage_path": res["path"], "visibility": "public", "content_type": ct,
                               "size": res.get("size"), "original_filename": file.filename, "is_deleted": False,
                               "uploaded_by": user["id"], "created_at": now_iso()})
    return {"url": f"/api/files/public/{res['path']}"}


# ---------------- promotions ----------------
class PromotionIn(BaseModel):
    title_en: str = Field(min_length=1, max_length=120)
    title_my: Optional[str] = ""
    subtitle_en: Optional[str] = ""
    subtitle_my: Optional[str] = ""
    image_url: Optional[str] = ""
    link: Optional[str] = "/products"
    placement: Literal["hero", "featured"] = "featured"
    start_at: Optional[str] = None
    end_at: Optional[str] = None
    active: bool = True
    sort: int = 100


@router.get("/admin/promotions")
async def list_promos(user=Depends(require("promotions"))):
    return await db.promotions.find({"deleted_at": None}, {"_id": 0}).sort("sort", 1).to_list(200)


@router.post("/admin/promotions")
async def create_promo(body: PromotionIn, user=Depends(require("promotions"))):
    d = {"id": new_id(), **body.model_dump(), "created_at": now_iso(), "deleted_at": None}
    await db.promotions.insert_one(dict(d))
    await audit(actor_of(user), "promotion.created", "promotion", d["id"], {"title": body.title_en})
    return d


@router.put("/admin/promotions/{iid}")
async def update_promo(iid: str, body: PromotionIn, user=Depends(require("promotions"))):
    await db.promotions.update_one({"id": iid}, {"$set": body.model_dump()})
    await audit(actor_of(user), "promotion.updated", "promotion", iid)
    return await db.promotions.find_one({"id": iid}, {"_id": 0})


@router.delete("/admin/promotions/{iid}")
async def delete_promo(iid: str, user=Depends(require("promotions"))):
    await db.promotions.update_one({"id": iid}, {"$set": {"deleted_at": now_iso()}})
    await audit(actor_of(user), "promotion.deleted", "promotion", iid)
    return {"ok": True}


# ---------------- payment methods ----------------
class PaymentMethodIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)
    account_name: Optional[str] = Field(default="", max_length=80)
    account_number: Optional[str] = Field(default="", max_length=40)
    instructions: Optional[str] = Field(default="", max_length=600)
    enabled: bool = False
    sort: int = 100


@router.get("/admin/payment-methods")
async def list_pm(user=Depends(require("orders"))):
    return await db.payment_methods.find({}, {"_id": 0}).sort("sort", 1).to_list(50)


@router.post("/admin/payment-methods")
async def create_pm(body: PaymentMethodIn, user=Depends(require("payment_methods"))):
    d = {"id": new_id(), "code": slugify(body.name).rsplit("-", 1)[0][:12] + new_id()[:4], **body.model_dump(), "enabled": False, "created_at": now_iso()}
    await db.payment_methods.insert_one(dict(d))
    await audit(actor_of(user), "payment_method.created", "payment_method", d["id"], {"name": body.name})
    return d


@router.put("/admin/payment-methods/{iid}")
async def update_pm(iid: str, body: PaymentMethodIn, user=Depends(require("payment_methods"))):
    if body.enabled and not (body.account_name and body.account_number):
        raise HTTPException(400, "Enter the real recipient account name and number before enabling this method")
    r = await db.payment_methods.update_one({"id": iid}, {"$set": body.model_dump()})
    if not r.matched_count:
        raise HTTPException(404, "Not found")
    await audit(actor_of(user), "payment_method.updated", "payment_method", iid, {"enabled": body.enabled, "name": body.name})
    return await db.payment_methods.find_one({"id": iid}, {"_id": 0})


# ---------------- settings & integrations ----------------
@router.get("/admin/settings")
async def get_settings(user=Depends(require("settings"))):
    return await db.settings.find_one({"id": "global"}, {"_id": 0})


@router.put("/admin/settings")
async def put_settings(body: dict, user=Depends(require("settings"))):
    allowed = {"brand_name", "logo_url", "og_image", "tagline_en", "tagline_my", "contacts", "messenger_url", "hero",
               "policies", "faq", "services"}
    upd = {k: v for k, v in body.items() if k in allowed}
    if isinstance(upd.get("faq"), list):
        upd["faq"] = [{k: str(q.get(k) or "") for k in ("q_en", "q_my", "a_en", "a_my")} for q in upd["faq"] if isinstance(q, dict)]
    await db.settings.update_one({"id": "global"}, {"$set": {**upd, "updated_at": now_iso()}})
    await audit(actor_of(user), "settings.updated", "settings", "global", {"fields": list(upd)})
    return await db.settings.find_one({"id": "global"}, {"_id": 0})


@router.get("/admin/integrations")
async def integrations(user=Depends(require("dashboard"))):
    tg = {"configured": tg_api.configured(), "bot_username": tg_api.bot_username() or None,
          "owner_configured": bool(tg_api.owner_id()), "webhook": None,
          "public_base_url": os.environ.get("PUBLIC_BASE_URL")}
    if tg_api.configured():
        ok, info = await tg_api.tg_call("getWebhookInfo")
        if ok:
            tg["webhook"] = {"url": info.get("url"), "pending_update_count": info.get("pending_update_count"),
                             "last_error_message": info.get("last_error_message"), "last_error_date": info.get("last_error_date")}
    s = await db.settings.find_one({"id": "global"}, {"_id": 0}) or {}
    msgr = {"mode": "click_to_messenger", "messenger_url": s.get("messenger_url"), "official_api_connector": "disabled",
            "checklist": ["Facebook Page with admin access", "Meta developer app with Messenger product",
                          "pages_messaging permission approved in App Review", "Page access token (stored as a server secret)",
                          "Webhook verify token + app secret for signature verification", "Privacy policy URL published"]}
    return {"telegram": tg, "messenger": msgr}


@router.post("/admin/integrations/telegram/webhook")
async def register_wh(user=Depends(require("settings"))):
    ok, res = await TB.register_webhook(os.environ.get("PUBLIC_BASE_URL"))
    await audit(actor_of(user), "telegram.webhook_registered", "integration", "telegram", {"ok": ok})
    if not ok:
        raise HTTPException(400, f"Webhook registration failed: {res}")
    return {"ok": True}


@router.post("/admin/integrations/telegram/test")
async def tg_test(user=Depends(require("settings"))):
    ok, res = await tg_api.notify_owner("✅ Test message from your store admin dashboard.")
    if not ok:
        raise HTTPException(400, f"Could not message the owner: {res}. The owner must press Start in the bot first.")
    return {"ok": True}


# ---------------- dashboard ----------------
@router.get("/admin/dashboard")
async def dashboard(days: int = 30, user=Depends(require("dashboard"))):
    since = (now() - timedelta(days=max(1, min(days, 365)))).isoformat()
    c = db.orders.count_documents
    by = lambda ss: c({"status": {"$in": ss}, "created_at": {"$gte": since}})  # noqa
    rev = await db.orders.aggregate([{"$match": {"status": "COMPLETED", "completed_at": {"$gte": since}}},
                                     {"$group": {"_id": {"$substr": ["$completed_at", 0, 10]}, "t": {"$sum": "$total_mmk"}, "n": {"$sum": 1}}},
                                     {"$sort": {"_id": 1}}]).to_list(400)
    sups = await db.suppliers.find({"deleted_at": None}, {"_id": 0, "id": 1, "name": 1, "health": 1, "last_sync_at": 1, "last_error": 1, "active": 1}).to_list(50)
    return {
        "products_total": await db.products.count_documents({"deleted_at": None}),
        "products_published": await db.products.count_documents({"deleted_at": None, "status": "published"}),
        "products_low_stock": await db.products.count_documents({"deleted_at": None, "status": "published", "$or": [{"low_stock": True}, {"available": False}]}),
        "orders": {
            "pending": await by(["CHECKOUT_STARTED", "AWAITING_CUSTOMER_CONFIRMATION"]),
            "awaiting_payment": await by(["AWAITING_PAYMENT"]),
            "payment_review": await by(["PAYMENT_SUBMITTED", "PAYMENT_UNDER_REVIEW"]),
            "paid": await by(["PAID_AWAITING_FULFILLMENT", "FULFILLMENT_IN_PROGRESS", "DELIVERED", "COMPLETED"]),
            "delivery_pending": await by(["PAID_AWAITING_FULFILLMENT", "FULFILLMENT_IN_PROGRESS"]),
            "completed": await by(["COMPLETED"]),
            "cancelled": await by(["CANCELLED"]),
        },
        "revenue_mmk": sum(r["t"] for r in rev),
        "revenue_series": [{"date": r["_id"], "revenue": r["t"], "orders": r["n"]} for r in rev],
        "suppliers": sups,
        "recent_activity": await db.audit_logs.find({}, {"_id": 0}).sort("created_at", -1).limit(12).to_list(12),
    }


@router.get("/admin/audit")
async def audit_logs(entity_type: str = "", limit: int = 200, user=Depends(require("audit"))):
    q = {"entity_type": entity_type} if entity_type else {}
    return await db.audit_logs.find(q, {"_id": 0}).sort("created_at", -1).limit(min(limit, 500)).to_list(500)


# ---------------- orders ----------------
async def get_order(oid):
    o = await db.orders.find_one({"$or": [{"id": oid}, {"ref": oid.upper()}]}, {"_id": 0, "track_hash": 0})
    if not o:
        raise HTTPException(404, "Order not found")
    return o


def decorate(o):
    o["payment_status"] = O.payment_status(o["status"])
    o["delivery_status"] = O.delivery_status(o["status"])
    o.pop("track_hash", None)
    return o


@router.get("/admin/orders")
async def list_orders(status: str = "", source: str = "", q: str = "", assignee: str = "", user=Depends(require("orders"))):
    import re
    query = {}
    if status:
        query["status"] = {"$in": status.split(",")}
    if source:
        query["source"] = source
    if assignee:
        query["assignee_id"] = user["id"] if assignee == "me" else assignee
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        query["$or"] = [{"ref": rx}, {"customer_name": rx}, {"customer_contact": rx}, {"items.product_name": rx}, {"messenger_ref": rx}]
    items = await db.orders.find(query, {"_id": 0}).sort("created_at", -1).limit(500).to_list(500)
    return [decorate(o) for o in items]


@router.get("/admin/orders/{oid}")
async def order_detail(oid: str, user=Depends(require("orders"))):
    o = decorate(await get_order(oid))
    o["events"] = await db.order_events.find({"order_id": o["id"]}, {"_id": 0}).sort("created_at", 1).to_list(200)
    o["evidence"] = await db.payment_evidence.find({"order_id": o["id"]}, {"_id": 0}).sort("created_at", 1).to_list(50)
    o["reviews"] = await db.payment_reviews.find({"order_id": o["id"]}, {"_id": 0}).to_list(50)
    o["deliveries"] = await db.deliveries.find({"order_id": o["id"]}, {"_id": 0}).to_list(20)
    o["artifacts"] = [{**{k: v for k, v in a.items() if k != "content_enc"}, "has_content": bool(a.get("content_enc"))}
                      for a in await db.delivery_artifacts.find({"order_id": o["id"]}, {"_id": 0}).to_list(20)]
    if o.get("payment_method_id"):
        o["payment_method"] = await db.payment_methods.find_one({"id": o["payment_method_id"]}, {"_id": 0, "name": 1})
    if o.get("assignee_id"):
        a = await db.users.find_one({"id": o["assignee_id"]}, {"_id": 0, "name": 1, "email": 1})
        o["assignee"] = a
    o["allowed_transitions"] = sorted(O.TRANSITIONS.get(o["status"], set()))
    return o


class ManualOrderIn(BaseModel):
    product_id: str
    variant_id: str
    source: Literal["ADMIN_MANUAL", "DIRECT_MESSENGER"] = "ADMIN_MANUAL"
    customer_name: str = Field(min_length=1, max_length=80)
    customer_contact: Optional[str] = Field(default="", max_length=120)
    messenger_ref: Optional[str] = Field(default="", max_length=80)
    note: Optional[str] = Field(default="", max_length=1000)


@router.post("/admin/orders")
async def manual_order(body: ManualOrderIn, user=Depends(require("orders"))):
    try:
        p, v = await O.load_product_variant(body.product_id, body.variant_id)
    except O.OrderError as e:
        raise HTTPException(400, str(e)) from e
    o = await O.create_order(p, v, body.source, actor_of(user), "AWAITING_PAYMENT",
                             channel="messenger" if body.source == "DIRECT_MESSENGER" else "manual",
                             customer_name=body.customer_name, customer_contact=body.customer_contact, note=body.note,
                             extra={"messenger_ref": body.messenger_ref or None, "assignee_id": user["id"]})
    return decorate(o)


class TransitionIn(BaseModel):
    to: str
    note: Optional[str] = Field(default="", max_length=1000)


@router.post("/admin/orders/{oid}/transition")
async def do_transition(oid: str, body: TransitionIn, user=Depends(require("orders"))):
    o = await get_order(oid)
    if body.to in O.APPROVAL_TARGETS and not has_perm(user, "orders.approve"):
        raise HTTPException(403, "You are not allowed to approve/reject payments or refunds")
    if body.to in O.DELIVERY_TARGETS and not has_perm(user, "orders.deliver"):
        raise HTTPException(403, "You are not allowed to manage delivery")
    try:
        if body.to == "PAID_AWAITING_FULFILLMENT":
            res, _ = await O.approve_payment(o["id"], actor_of(user), body.note)
        elif body.to == "PAYMENT_REJECTED":
            res, _ = await O.reject_payment(o["id"], actor_of(user), body.note)
        elif body.to == "COMPLETED":
            res, _ = await O.complete(o["id"], actor_of(user), body.note)
        elif body.to == "DELIVERED":
            res, _ = await O.confirm_manual_delivery(o["id"], actor_of(user), body.note)
        elif body.to == "AWAITING_PAYMENT" and o["status"] in {"PAYMENT_SUBMITTED", "PAYMENT_UNDER_REVIEW"}:
            res, _ = await O.request_more_evidence(o["id"], actor_of(user), body.note)
        else:
            res, _ = await O.transition(o["id"], body.to, actor_of(user), note=body.note)
    except O.OrderError as e:
        raise HTTPException(400, str(e)) from e
    return decorate(res)


class AssignIn(BaseModel):
    user_id: Optional[str] = None
    note: Optional[str] = None


@router.post("/admin/orders/{oid}/assign")
async def assign(oid: str, body: AssignIn, user=Depends(require("orders"))):
    o = await get_order(oid)
    upd = {"assignee_id": body.user_id}
    if body.note is not None:
        upd["notes"] = body.note[:1000]
    await db.orders.update_one({"id": o["id"]}, {"$set": upd})
    await O.add_event(o["id"], o["status"], o["status"], actor_of(user), "assigned", body.user_id)
    await audit(actor_of(user), "order.assigned", "order", o["id"], {"assignee": body.user_id})
    return {"ok": True}


@router.post("/admin/orders/{oid}/evidence")
async def add_evidence(oid: str, reference: str = Form(""), file: Optional[UploadFile] = File(None), user=Depends(require("orders"))):
    o = await get_order(oid)
    if o["status"] not in {"AWAITING_PAYMENT", "PAYMENT_REJECTED", "PAYMENT_SUBMITTED", "PAYMENT_UNDER_REVIEW"}:
        raise HTTPException(400, f"Order is {o['status']}; evidence not accepted")
    ev = {"id": new_id(), "order_id": o["id"], "kind": "transaction_reference", "reference": reference[:200] or None,
          "storage_path": None, "content_type": None, "review_status": "pending", "submitted_via": "admin", "submitted_by": user["id"], "created_at": now_iso()}
    if file is not None and file.filename:
        ct = file.content_type or ""
        if ct not in IMAGE_TYPES:
            raise HTTPException(400, "Evidence must be an image")
        data = await file.read()
        if len(data) > 8 * 1024 * 1024:
            raise HTTPException(400, "Image too large")
        res = await put_object(f"{APP_NAME}/private/evidence/{o['id']}/{ev['id']}.{IMAGE_TYPES[ct]}", data, ct)
        ev.update(kind="screenshot", storage_path=res["path"], content_type=ct)
    elif not reference:
        raise HTTPException(400, "Provide a transaction reference or screenshot")
    await db.payment_evidence.insert_one(dict(ev))
    if o["status"] in {"AWAITING_PAYMENT", "PAYMENT_REJECTED"}:
        await O.transition(o["id"], "PAYMENT_SUBMITTED", actor_of(user), "payment_evidence_submitted", ev["reference"] or ev["kind"])
    await tg_api.notify_owner("🔔 <b>Payment review needed</b> (added by staff)\n" + TB.order_summary(await get_order(o["id"]), True),
                              TB.admin_buttons(await get_order(o["id"])))
    ev.pop("_id", None)
    return ev


@router.get("/admin/evidence/{eid}/file")
async def evidence_file(eid: str, user=Depends(require("orders"))):
    ev = await db.payment_evidence.find_one({"id": eid}, {"_id": 0})
    if not ev or not ev.get("storage_path"):
        raise HTTPException(404, "No file")
    data, ct = await get_object(ev["storage_path"])
    await audit(actor_of(user), "evidence.viewed", "order", ev["order_id"], {"evidence_id": eid})
    return Response(content=data, media_type=ev.get("content_type") or ct, headers={"Cache-Control": "private, no-store"})


class DeliverIn(BaseModel):
    artifact_type: Literal["license", "activation_link", "instructions", "credentials", "other"] = "instructions"
    content: str = Field(min_length=1, max_length=4000)


@router.post("/admin/orders/{oid}/deliver")
async def deliver(oid: str, body: DeliverIn, user=Depends(require("orders.deliver"))):
    o = await get_order(oid)
    try:
        res = await O.deliver(o["id"], body.artifact_type, body.content, actor_of(user))
    except O.OrderError as e:
        raise HTTPException(400, str(e)) from e
    return {"result": res["result"], "order": decorate(res["order"])}


@router.get("/admin/orders/{oid}/artifacts/{aid}")
async def reveal_artifact(oid: str, aid: str, user=Depends(require("orders.deliver"))):
    a = await db.delivery_artifacts.find_one({"id": aid}, {"_id": 0})
    o = await get_order(oid)
    if not a or a["order_id"] != o["id"]:
        raise HTTPException(404, "Not found")
    if not a.get("content_enc"):
        raise HTTPException(410, "Artifact content was purged after completion")
    await audit(actor_of(user), "artifact.revealed", "order", a["order_id"], {"artifact_id": aid})
    return {"content": decrypt(a["content_enc"]), "type": a["type"]}
